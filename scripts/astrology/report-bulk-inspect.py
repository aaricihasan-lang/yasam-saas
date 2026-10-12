# 2026-10-12 UAT 2. tur (madde L) — report-bulk.e2e.mjs çıktılarını AÇAR ve içeriği denetler (python-docx + PyMuPDF).
# Çalıştırma: python scripts/astrology/report-bulk-inspect.py <out-dir> [png-dir]
import json
import os
import sys

import docx
import pymupdf

OUT = sys.argv[1]
PNG = sys.argv[2] if len(sys.argv) > 2 else None
ok = 0
bad = 0


def t(name, cond, extra=""):
    global ok, bad
    if cond:
        ok += 1
    else:
        bad += 1
    print(("PASS " if cond else "FAIL ") + name + (" — " + extra if extra else ""))


meta = json.load(open(os.path.join(OUT, "meta.json"), encoding="utf-8"))
client = meta["profileA"]
SECTIONS = {
    "1-tekli-transit": ["Transit"],
    "2-secili-3": ["Natal", "Transit", "Sinastri"],
    "3-tumu-5": ["Natal", "Transit", "Progression", "Sinastri", "Forecast"],
}
# Çark beklenen: Natal 1, Transit 1 (çift halka), Progression 1, Sinastri 1 (çift halka), Forecast 0
CHARTS = {"1-tekli-transit": 1, "2-secili-3": 3, "3-tumu-5": 4}


def techniques_in(text):
    keys = {
        "Natal": ["Natal harita", "Doğum haritası", "Natal"],
        "Transit": ["Transit"],
        "Progression": ["İlerletim", "Progression", "progression"],
        "Sinastri": ["Sinastri", "Synastry"],
        "Forecast": ["Forecast", "Takvim", "Transit takvimi"],
    }
    return {k for k, vs in keys.items() if any(v in text for v in vs)}


for base, techs in SECTIONS.items():
    # ── Word
    d = docx.Document(os.path.join(OUT, base + ".docx"))
    paras = [p.text for p in d.paragraphs]
    body = "\n".join(paras)
    tables = d.tables
    cells = "\n".join(c.text for tb in tables for row in tb.rows for c in row.cells)
    full = body + "\n" + cells
    imgs = [r for r in d.part.rels.values() if "image" in r.reltype]
    def style_id(p):
        ps = p._p.pPr.pStyle if p._p.pPr is not None else None
        return ps.val if ps is not None else ""

    heads = [p.text for p in d.paragraphs if style_id(p).startswith("Heading")]
    t(f"{base}.docx açıldı", len(paras) > 20, f"{len(paras)} paragraf, {len(tables)} tablo, {len(imgs)} görsel")
    t(f"{base}.docx danışan adı + Türkçe karakter", client in full and all(ch in full for ch in "ÇŞÖİçşğıöü"))
    t(f"{base}.docx doğum bilgisi revizyonu", "revizyon" in full.lower())
    t(f"{base}.docx gerçek çark görseli sayısı", len(imgs) >= CHARTS[base], f"{len(imgs)} görsel (beklenen ≥ {CHARTS[base]}; logo/kapak dahil olabilir)")
    t(f"{base}.docx tablo başlık satırı tekrarı (tblHeader)", sum(1 for tb in tables if tb.rows and tb.rows[0]._tr.trPr is not None and tb.rows[0]._tr.trPr.xml.find("tblHeader") >= 0) >= 2)
    if len(techs) > 1:
        t(f"{base}.docx İçindekiler", "İçindekiler" in heads)
        h1s = [p.text for p in d.paragraphs if style_id(p) == "Heading1"]
        sec = [h for h in h1s if h.split(". ")[0].isdigit()]
        t(f"{base}.docx her analiz ayrı bölüm (Heading 1)", len(sec) == len(techs), " | ".join(sec))
        t(f"{base}.docx bölümler yeni sayfada", sum(1 for p in d.paragraphs if p.paragraph_format.page_break_before) >= len(techs))
    t(f"{base}.docx uzman notları + kayıt bilgisi bölümleri", "Uzman notları" in heads and "Kayıt bilgisi" in heads)
    t(f"{base}.docx '[object Object]' / 'undefined' yok", "[object Object]" not in full and "undefined" not in full)

    # ── PDF
    pdf = pymupdf.open(os.path.join(OUT, base + ".pdf"))
    ptext = "\n".join(pg.get_text() for pg in pdf)
    pimgs = len({im[0] for pg in pdf for im in pg.get_images()})
    t(f"{base}.pdf açıldı", pdf.page_count >= 3, f"{pdf.page_count} sayfa, {pimgs} görsel")
    t(f"{base}.pdf Türkçe karakter + danışan", client in ptext.replace("\n", " ") or client.split(" ")[1] in ptext, "")
    t(f"{base}.pdf çark görseli", pimgs >= CHARTS[base], f"{pimgs} (beklenen ≥ {CHARTS[base]})")
    t(f"{base}.pdf sayfa numarası", f"Sayfa 1 / {pdf.page_count}" in ptext)
    if len(techs) > 1:
        toc = pdf.get_toc()
        t(f"{base}.pdf yer imleri (her analiz)", len(toc) == len(techs), " | ".join(f"{x[1]}→s{x[2]}" for x in toc))
        t(f"{base}.pdf İçindekiler", "İçindekiler" in ptext)
    t(f"{base}.pdf '[object Object]' / 'undefined' yok", "[object Object]" not in ptext and "undefined" not in ptext)
    if PNG:
        os.makedirs(PNG, exist_ok=True)
        for i, pg in enumerate(pdf):
            if pg.get_images():
                pg.get_pixmap(dpi=70).save(os.path.join(PNG, f"{base}-s{i + 1}.png"))

print(f"\nreport-bulk inspect: {ok} PASS / {bad} FAIL")
sys.exit(1 if bad else 0)
